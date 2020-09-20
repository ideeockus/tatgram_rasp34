from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy import create_engine, MetaData, Table, Column, Integer, String, Boolean
from sqlalchemy.ext.declarative import declarative_base
from sqlalchemy.orm import sessionmaker
import datetime


"""
A - 0 empty column
B - 1 class_name
C - 2 week_day
D -3 lesson_start_time
E -4 lesson_end_time
F - 5 subject_name
G - 6 room_number
"""

# speed test
script_start_time = datetime.datetime.now()

Base = declarative_base()  # декларативный базовый класс


class Lessons(Base):
    __tablename__ = "rasp_db"
    id = Column(Integer, primary_key=True)
    class_name = Column(String)
    week_day = Column(String)
    lesson_start_time = Column(String)
    lesson_end_time = Column(String)
    subject_name = Column(String)
    room_number = Column(String)


postgres_db = "<ЗДЕСЬ ДОЛЖНА БЫТЬ ССЫЛКА С АВТОРИЗАЦИЕЙ В БД>"  # для запуска нужно вписать ссылку со всеми данными к БД
# engine = create_engine('sqlite:///databases/rasp.db', echo=True)
engine = create_engine(postgres_db, echo=False)
Base.metadata.create_all(engine)

Session = sessionmaker(bind=engine)
rasp_session = Session()

rasp_excel = load_workbook("databases/raspisanie.xlsx", read_only=True)
rasp = rasp_excel.active

for row_num in range(3, rasp.max_row):
    class_name = rasp[row_num][1].value
    week_day = rasp[row_num][2].value
    lesson_start_time = str(rasp[row_num][3].value)
    lesson_end_time = str(rasp[row_num][4].value)
    subject_name = rasp[row_num][5].value
    room_number = rasp[row_num][6].value

    if class_name is None:
        continue

    lesson = Lessons(
        class_name=class_name,
        week_day=week_day,
        lesson_start_time=lesson_start_time,
        lesson_end_time=lesson_end_time,
        subject_name=subject_name,
        room_number=room_number,
    )
    rasp_session.add(lesson)
    print(class_name, week_day, lesson_start_time, subject_name, room_number)

rasp_session.commit()

# speed test
script_end_time = datetime.datetime.now()
print((script_end_time-script_start_time).total_seconds())
